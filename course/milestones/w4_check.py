"""
course/milestones/w4_check.py
4-hafta milestone (M4) o'z-o'zini tekshiruvchi skript — kursning yakuniy tekshiruvi.

Uch vazifani tekshiradi:
    A) SentimentAPI (capstone/app.py): TestClient orqali POST /predict ->
       {"sentiment": ijobiy|salbiy, "confidence": float in [0,1]}.
    B) Bilim testi: course/final_test.docx (>=30 savol) + final_test.xlsx (javoblar kaliti).
    C) Agent: m15 import + DocumentAssistantAgent().run() -> str.

Ishga tushirish (repo ildizidan yoki istalgan joydan):
    python course/milestones/w4_check.py

Vositalar mahalliy bor: fastapi/httpx (TestClient), python-docx (docx), openpyxl.
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent          # course/milestones
REPO = HERE.parent.parent                        # repo ildizi
sys.path.insert(0, str(REPO / "capstone"))
sys.path.insert(0, str(REPO / "capstone" / "modules"))

_passed = 0


def check(nom: str, shart: bool) -> None:
    global _passed
    assert shart, f"FAIL: {nom}"
    print(f"  ✓ {nom}")
    _passed += 1


print("=" * 70)
print("4-HAFTA MILESTONE (M4) — O'Z-O'ZINI TEKSHIRUV (w4_check)")
print("Deploy (FastAPI) + bilim testi (L1-L14) + agent (m15)")
print("=" * 70)

# ══════════════════════════════════════════════════════════════════════════
# TASK A — SentimentAPI: POST /predict
# ══════════════════════════════════════════════════════════════════════════
print("\n[Task A] SentimentAPI — POST /predict (FastAPI TestClient)")

from fastapi.testclient import TestClient
import app as sentiment_app

check("app.py: create_sentiment_api() mavjud", callable(getattr(sentiment_app, "create_sentiment_api", None)))
check("app.py: modul darajasida `app` (FastAPI) mavjud", getattr(sentiment_app, "app", None) is not None)

client = TestClient(sentiment_app.app)

r_pos = client.post("/predict", json={"text": "mahsulot juda sifatli va arzon"})
check("POST /predict status 200", r_pos.status_code == 200)
body = r_pos.json()
check("javobda 'sentiment' va 'confidence' kalitlari bor",
      set(body.keys()) >= {"sentiment", "confidence"})
check("sentiment ijobiy/salbiy (qulflangan yorliqlar)", body["sentiment"] in ("ijobiy", "salbiy"))
check("confidence float [0,1]", isinstance(body["confidence"], (int, float)) and 0.0 <= body["confidence"] <= 1.0)

r_neg = client.post("/predict", json={"text": "mahsulot buzuq keldi juda yomon"})
check("salbiy sharh ham ijobiy/salbiy qaytaradi", r_neg.json()["sentiment"] in ("ijobiy", "salbiy"))
print(f"    /predict('...sifatli...') -> {body}")

# ══════════════════════════════════════════════════════════════════════════
# TASK B — Bilim testi: final_test.docx + final_test.xlsx
# ══════════════════════════════════════════════════════════════════════════
print("\n[Task B] Bilim testi — final_test.docx + final_test.xlsx (L1-L14)")

DOCX = REPO / "course" / "final_test.docx"
XLSX = REPO / "course" / "final_test.xlsx"
check("final_test.docx mavjud", DOCX.exists())
check("final_test.xlsx mavjud", XLSX.exists())

import docx  # python-docx
doc = docx.Document(str(DOCX))
paras = [p.text for p in doc.paragraphs if p.text.strip()]
n_savol = sum(1 for t in paras if t.lstrip().split(".", 1)[0].strip().isdigit())
check("docx ochiladi va >=30 raqamlangan savol bor", n_savol >= 30)

import openpyxl
wb = openpyxl.load_workbook(str(XLSX))
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
check("xlsx ochiladi va javoblar kaliti to'liq (>=31 qator, sarlavha + 30)", len(rows) >= 31)
print(f"    docx savollar: {n_savol} | xlsx qatorlar: {len(rows)}")

# ══════════════════════════════════════════════════════════════════════════
# TASK C — Agent: m15 import + run()
# ══════════════════════════════════════════════════════════════════════════
print("\n[Task C] Agent — m15 DocumentAssistantAgent (P15 da qurilgan)")

import m15_langchain_agent as m15
m15.USE_LANGCHAIN = False
check("m15: DocumentAssistantAgent mavjud", hasattr(m15, "DocumentAssistantAgent"))
agent = m15.DocumentAssistantAgent(tools={"sentiment_classify": lambda t: "ijobiy"})
out = agent.run("Bu sharh ijobiymi?")
check("m15 run() bo'sh bo'lmagan str qaytaradi", isinstance(out, str) and bool(out))
check("m15 route() L15 [I1]: sentiment_classify, conf>0.7",
      agent.route("Uzum Market ilovasi yaxshimi?")["tool"] == "sentiment_classify"
      and agent.route("Uzum Market ilovasi yaxshimi?")["confidence"] > 0.7)

# ══════════════════════════════════════════════════════════════════════════
print("\n" + "=" * 70)
print(f"NATIJA: {_passed} tekshiruvning hammasi O'TDI ✓")
print("4-hafta milestone tayyor: deploy + bilim testi + agent. KURS YAKUNLANDI. \U0001F393")
print("=" * 70)
